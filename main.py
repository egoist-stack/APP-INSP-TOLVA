import flet as ft
import pandas as pd
import openpyxl
from openpyxl.worksheet.pagebreak import Break
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.worksheet.properties import PageSetupProperties
import os
import io
import asyncio
import tempfile
import math
import zipfile
import re
import json
from PIL import Image, ImageOps, ImageDraw
import flet.canvas as cv
from datetime import date

try:  # opcional: si el APK incluye pillow-heif, también se leen fotos .heic
    import pillow_heif
    pillow_heif.register_heif_opener()
except Exception:
    pass

# ==============================================================================
# 1. BASE DE DATOS LOCAL Y CONSTANTES FIELES AL ORIGINAL
# ==============================================================================
def _bases_busqueda():
    bases = []
    try:
        bases.append(os.path.dirname(os.path.abspath(__file__)))
    except NameError:
        pass
    bases.append(os.getcwd())
    for var in ("FLET_APP_DIR", "FLET_ASSETS_DIR"):
        v = os.environ.get(var)
        if v:
            bases.append(v)
    vistas, unicas = set(), []
    for b in bases:
        if b not in vistas:
            vistas.add(b)
            unicas.append(b)
    return unicas

def buscar_recurso(*partes):
    """Busca un archivo o carpeta junto a main.py, en el cwd o en assets/ (funciona en PC y en el APK)."""
    for base in _bases_busqueda():
        for sub in ("", "assets"):
            ruta = os.path.join(base, sub, *partes)
            if os.path.exists(ruta):
                return ruta
    return None

def diagnostico_archivos():
    lineas = []
    for base in _bases_busqueda()[:2]:
        try:
            lineas.append(f"{base}: " + ", ".join(sorted(os.listdir(base))[:20]))
        except Exception as ex:
            lineas.append(f"{base}: no se pudo leer ({ex})")
    return "\n".join(lineas)

DB_FILE = os.path.join(os.environ.get("FLET_APP_STORAGE_DATA") or ".", "base_datos", "tolvas_db.json")

def cargar_db():
    if os.path.exists(DB_FILE):
        with open(DB_FILE, "r", encoding="utf-8") as f:
            try:
                return json.load(f)
            except Exception:
                return {}
    return {}

def guardar_db(data):
    try:
        os.makedirs(os.path.dirname(DB_FILE), exist_ok=True)
        with open(DB_FILE, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=4)
    except Exception:
        pass

ROTAR_FOTOS_VERTICALES = False

_NOMBRE_COMPLETO_DEFECTO = {
    "D": "DESPRENDIMIENTO",
    "DE": "DESGASTE",
    "DP": "DESPRENDIMIENTO PARCIAL",
    "F": "FISURA",
    "FA": "FRACTURA",
    "LF": "LIBRE DE FISURAS",
}

ESTRUCTURA_ZONAS = [
    {
        "titulo": "ZONA 01: CONJUNTO DE BLINDAJE DE TOLVA",
        "esquema": ["CONJUNTO DE BLINDAJE DE TOLVA.png"],
        "items": [
            ("1.1", "REFUERZO DE PISO #1", "VT/UT"),
            ("1.2", "REFUERZO DE PISO #2", "VT/UT"),
            ("1.3", "REFUERZO DE PISO #3", "VT/UT"),
            ("1.4", "ROCKBOX", "VT"),
            ("1.5", "REFUERZO FRONTAL #1", "VT/UT"),
            ("1.6", "REFUERZO FRONTAL #2", "VT/UT"),
            ("1.7", "REFUERZO FRONTAL #3", "VT/UT"),
            ("1.8", "REFUERZO LATERAL RH", "VT/UT"),
            ("1.9", "REFUERZO LATERAL LH", "VT"),
            ("1.10", "CORTAFLUJOS", "VT/UT")
        ]
    },
    {
        "titulo": "ZONA 02: CONJUNTO LATERAL (RH / LH)",
        "esquema": ["CONJUNTO LATERAL.png"],
        "items": [
            ("2.1", "PLANCHA LATERAL RH", "VT"),
            ("2.2", "VIGA CAJÓN LATERAL RH", "VT"),
            ("2.3", "PLANCHA LATERAL LH", "VT"),
            ("2.4", "VIGA CAJÓN LATERAL LH", "VT")
        ]
    },
    {
        "titulo": "ZONA 03: CANOPY",
        "esquema": ["CANOPY.png"],
        "items": [
            ("3.1", "PLANCHA RH", "VT"),
            ("3.2", "PLANCHA LH", "VT"),
            ("3.3", "DEFLECTOR RH", "VT"),
            ("3.4", "DEFLECTOR LH", "VT"),
            ("3.5", "PLANCHA FRONTAL CANOPY", "VT"),
            ("3.6", "CARTELAS DE PLANCHA FRONTAL", "VT"),
            ("3.7", "VIGA LATERAL RH", "VT"),
            ("3.8", "VIGA LATERAL LH", "VT"),
            ("3.9", "REFUERZO RH DE CANOPY", "VT"),
            ("3.10", "REFUERZO LH DE CANOPY", "VT")
        ]
    },
    {
        "titulo": "ZONA 04 Y 05: PLANCHAS DE PISO / PLANCHAS FRONTALES",
        "esquema": ["PLANCHAS FRONTALES.png", "PLANCHAS DE PISO.png"],
        "items": [
            ("4.1", "PLANCHA FRONTAL SUPERIOR", "VT"),
            ("4.2", "PLANCHA FRONTAL RH", "VT"),
            ("4.3", "PLANCHA FRONTAL LH", "VT"),
            ("5.1", "PLANCHA DE PISO RH", "VT/UT"),
            ("5.2", "PLANCHA DE PISO LH", "VT/UT"),
            ("5.3", "PLANCHA COLA DE PISO", "VT/UT")
        ]
    },
    {
        "titulo": "ZONA 06 Y 07: LONGUERINA DELANTERA, POSTERIOR Y VIGAS / GUIADORES",
        "esquema": ["LONGUERINA DELANTERA, POSTERIOR Y VIGAS.png", "GUIADORES.png"],
        "items": [
            ("6.1", "LONGUERINA DELANTERA RH", "VT"),
            ("6.2", "LONGUERINA DELANTERA LH", "VT"),
            ("6.3", "LONGUERINA POSTERIOR RH", "VT"),
            ("6.4", "LONGUERINA POSTERIOR LH", "VT"),
            ("6.5", "VIGA DE PISO #1 RH", "VT"),
            ("6.6", "VIGA CAJÓN DE PISO #1 LH", "VT"),
            ("6.7", "VIGA CAJÓN DE PISO #2 RH", "VT"),
            ("6.8", "VIGA CAJÓN DE PISO #2 LH", "VT"),
            ("6.9", "VIGA CAJÓN DE COLA", "VT"),
            ("7.1", "GUIADOR RH", "VT"),
            ("7.2", "GUIADOR LH", "VT")
        ]
    },
    {
        "titulo": "ZONA 08: CAJAS PIVOTE",
        "esquema": ["CAJAS PIVOTE.png"],
        "items": [
            ("8.1", "CAJA PIVOTE RH", "VT"),
            ("8.2", "CAJA PIVOTE LH", "VT"),
            ("8.3", "BUSHING DE CAJA PIVOTE RH", "VT"),
            ("8.4", "BUSHING DE CAJA PIVOTE LH", "VT"),
            ("8.5", "SEPARADOR DE CAJA PIVOTE", "VT")
        ]
    }
]

def es_pm_espesores(pm_str):
    if not pm_str:
        return False
    limpio = str(pm_str).strip().upper().replace("H", "").strip()
    try:
        val = int(limpio)
        return val > 0 and (val % 1000 == 0)
    except Exception:
        return False

# ==============================================================================
# 2. MOTOR DE PROCESAMIENTO DE EXCEL Y XML (100% FIEL AL ORIGINAL)
# ==============================================================================
def _insertar_filas_seguro(ws, fila_insercion, cantidad):
    merges_originales = list(ws.merged_cells.ranges)
    for mc in merges_originales:
        ws.unmerge_cells(str(mc))
    ws.insert_rows(fila_insercion, amount=cantidad)
    for mc in merges_originales:
        min_row, min_col, max_row, max_col = mc.min_row, mc.min_col, mc.max_row, mc.max_col
        if min_row >= fila_insercion:
            min_row += cantidad
            max_row += cantidad
        elif max_row >= fila_insercion:
            max_row += cantidad
        ws.merge_cells(start_row=min_row, start_column=min_col, end_row=max_row, end_column=max_col)

def _detectar_headers_zonas(ws):
    filas = []
    for row in range(1, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == "ZONA" and ws.cell(row=row, column=2).value == "DESCRIPCION":
            filas.append(row)
    return filas

def _crear_header_fotos(ws, fila_insercion):
    _insertar_filas_seguro(ws, fila_insercion, 1)
    ws.row_dimensions[fila_insercion].height = 25
    ws.merge_cells(start_row=fila_insercion, start_column=1, end_row=fila_insercion, end_column=4)
    ws.merge_cells(start_row=fila_insercion, start_column=5, end_row=fila_insercion, end_column=11)
    ws.merge_cells(start_row=fila_insercion, start_column=12, end_row=fila_insercion, end_column=18)
    borde = Border(left=Side(style="thick"), right=Side(style="thick"), top=Side(style="thick"), bottom=Side(style="thick"))
    align = Alignment(horizontal="center", vertical="center")
    for col, texto in [(1, "DESCRIPCION"), (5, "PANORAMICO"), (12, "DETALLE")]:
        c = ws.cell(row=fila_insercion, column=col, value=texto)
        c.font = Font(bold=True)
        c.alignment = align
    for col in range(1, 19):
        ws.cell(row=fila_insercion, column=col).border = borde

def _crear_bloque_foto_nuevo(ws, fila_insercion, alto=8):
    _insertar_filas_seguro(ws, fila_insercion, alto)
    for r in range(fila_insercion, fila_insercion + alto):
        ws.row_dimensions[r].height = 20
    ws.merge_cells(start_row=fila_insercion, start_column=1, end_row=fila_insercion + alto - 1, end_column=4)
    ws.merge_cells(start_row=fila_insercion, start_column=5, end_row=fila_insercion + alto - 1, end_column=11)
    ws.merge_cells(start_row=fila_insercion, start_column=12, end_row=fila_insercion + alto - 1, end_column=18)
    borde = Border(left=Side(style="thick"), right=Side(style="thick"), top=Side(style="thick"), bottom=Side(style="thick"))
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col in (1, 5, 12):
        c = ws.cell(row=fila_insercion, column=col)
        c.alignment = align
    for r in range(fila_insercion, fila_insercion + alto):
        for col in range(1, 19):
            ws.cell(row=r, column=col).border = borde
    return (fila_insercion, fila_insercion + alto - 1)

def _preparar_imagen_para_insertar(imagen_pil):
    img_copia = ImageOps.exif_transpose(imagen_pil).copy().convert("RGB")
    if ROTAR_FOTOS_VERTICALES and img_copia.height > img_copia.width:
        img_copia = img_copia.transpose(Image.ROTATE_90)
    img_copia.thumbnail((1200, 1200), Image.LANCZOS)
    buf = io.BytesIO()
    img_copia.save(buf, format="PNG")
    return buf.getvalue()

def _calcular_presupuesto_pagina_pt(ws):
    tamanos_papel_pt = {1: (612, 792), 9: (595, 842)}
    ancho_pt, alto_pt = tamanos_papel_pt.get(ws.page_setup.paperSize, (595, 842))
    if ws.page_setup.orientation == "landscape":
        ancho_pt, alto_pt = alto_pt, ancho_pt
    margenes = ws.page_margins
    top_pt = (margenes.top or 0.75) * 72
    bottom_pt = (margenes.bottom or 0.75) * 72
    imprimible_pt = alto_pt - top_pt - bottom_pt
    escala = (ws.page_setup.scale or 100) / 100
    return imprimible_pt / escala

def _altura_filas_pt(ws, fila_ini, fila_fin):
    if fila_fin < fila_ini:
        return 0
    total = 0
    for f in range(fila_ini, fila_fin + 1):
        dim = ws.row_dimensions.get(f)
        total += dim.height if (dim and dim.height) else 15
    return total

def _desplazar_filas_drawing_xml(xml_texto, fila_insercion_0idx, cantidad, altura_insertada_emu=0):
    def reemplazar_row(m):
        fila = int(m.group(2))
        if fila >= fila_insercion_0idx:
            fila += cantidad
        return f"{m.group(1)}{fila}{m.group(3)}"
    return re.sub(r'(<[a-zA-Z0-9]*:?row>)(\d+)(</[a-zA-Z0-9]*:?row>)', reemplazar_row, xml_texto)

def _ajustar_altura_forma_por_texto(xml_texto, texto_objetivo, nueva_altura_emu):
    def procesar_anchor(m):
        bloque = m.group(0)
        m_texto = re.search(r'<a:t>([^<]*)</a:t>', bloque)
        if not m_texto or m_texto.group(1).strip() != texto_objetivo:
            return bloque
        if '<xdr:grpSp>' in bloque:
            return bloque
        m_from_off = re.search(r'<xdr:from>.*?<xdr:rowOff>(\d+)</xdr:rowOff></xdr:from>', bloque, re.DOTALL)
        if not m_from_off:
            return bloque
        from_rowoff = int(m_from_off.group(1))
        nuevo_to_rowoff = from_rowoff + nueva_altura_emu
        bloque_nuevo = re.sub(
            r'(<xdr:to>.*?<xdr:rowOff>)\d+(</xdr:rowOff></xdr:to>)',
            lambda mm: f"{mm.group(1)}{nuevo_to_rowoff}{mm.group(2)}",
            bloque, count=1, flags=re.DOTALL
        )
        bloque_nuevo = re.sub(
            r'(<a:ext cx="\d+" cy=")\d+(")',
            lambda mm: f"{mm.group(1)}{nueva_altura_emu}{mm.group(2)}",
            bloque_nuevo, count=1
        )
        return bloque_nuevo
    return re.sub(r'<xdr:twoCellAnchor.*?</xdr:twoCellAnchor>', procesar_anchor, xml_texto, flags=re.DOTALL)

def _estandarizar_cajas_de_puntos(xml_texto, nuevo_cx_emu, nuevo_cy_emu):
    def procesar_shape(m):
        bloque = m.group(0)
        m_texto = re.search(r'<a:t>([^<]*)</a:t>', bloque)
        if not m_texto or not re.match(r'^\d+\.\d+$', m_texto.group(1).strip()):
            return bloque
        bloque_nuevo = re.sub(
            r'<a:ext cx="\d+" cy="\d+"/>',
            f'<a:ext cx="{nuevo_cx_emu}" cy="{nuevo_cy_emu}"/>',
            bloque, count=1
        )
        return bloque_nuevo
    return re.sub(r'<xdr:sp\s.*?</xdr:sp>', procesar_shape, xml_texto, flags=re.DOTALL)

def _quitar_anchors_en_ventana(xml_texto, fila_ini_0idx, fila_fin_0idx):
    anchors = re.findall(r'<xdr:twoCellAnchor.*?</xdr:twoCellAnchor>', xml_texto, re.DOTALL)
    resultado = xml_texto
    for a in anchors:
        m_from = re.search(r'<xdr:from><xdr:col>\d+</xdr:col>.*?<xdr:row>(\d+)</xdr:row>', a, re.DOTALL)
        if m_from and fila_ini_0idx <= int(m_from.group(1)) <= fila_fin_0idx:
            resultado = resultado.replace(a, '', 1)
    return resultado

class _Geometria:
    def __init__(self, ws, n_cols=18):
        self.ws = ws
        defecto = ws.sheet_format.defaultColWidth or 8.43
        self.cols = []
        for c in range(1, n_cols + 1):
            w = defecto
            for d in ws.column_dimensions.values():
                if d.min is not None and d.max is not None and d.min <= c <= d.max and d.width:
                    w = d.width
                    break
            px = int((256 * w + int(128 / 7)) / 256 * 7)
            self.cols.append(px * 9525)
        self.alto_defecto = ws.sheet_format.defaultRowHeight or 15

    def _alto_fila(self, r0):
        d = self.ws.row_dimensions.get(r0 + 1)
        h = d.height if (d is not None and d.height) else self.alto_defecto
        return int(h * 12700)

    def x(self, col0):
        return sum(self.cols[:col0])

    def y(self, fila0):
        return sum(self._alto_fila(r) for r in range(fila0))

    def col_en(self, x):
        acc = 0
        for i, w in enumerate(self.cols):
            if x < acc + w:
                return i, int(x - acc)
            acc += w
        return len(self.cols) - 1, self.cols[-1] - 1

    def fila_en(self, y):
        acc, r = 0, 0
        while True:
            h = self._alto_fila(r)
            if y < acc + h:
                return r, int(y - acc)
            acc += h
            r += 1

def _anchor_imagen_ajustada_xml(geo, id_imagen, nombre, rid, bytes_img,
                                col_ini, col_fin, fila_ini, fila_fin, margen=60000):
    iw, ih = Image.open(io.BytesIO(bytes_img)).size
    x0, x1 = geo.x(col_ini), geo.x(col_fin)
    y0, y1 = geo.y(fila_ini), geo.y(fila_fin)
    ancho = x1 - x0 - 2 * margen
    alto = y1 - y0 - 2 * margen
    if ancho <= 0 or alto <= 0 or iw <= 0 or ih <= 0:
        return ""
    esc = min(ancho / iw, alto / ih)
    cx, cy = int(iw * esc), int(ih * esc)
    px = x0 + margen + (ancho - cx) // 2
    py = y0 + margen + (alto - cy) // 2
    c, c_off = geo.col_en(px)
    r, r_off = geo.fila_en(py)
    return (
        f'<xdr:oneCellAnchor>'
        f'<xdr:from><xdr:col>{c}</xdr:col><xdr:colOff>{c_off}</xdr:colOff>'
        f'<xdr:row>{r}</xdr:row><xdr:rowOff>{r_off}</xdr:rowOff></xdr:from>'
        f'<xdr:ext cx="{cx}" cy="{cy}"/>'
        f'<xdr:pic>'
        f'<xdr:nvPicPr><xdr:cNvPr id="{id_imagen}" name="{nombre}{id_imagen}"/>'
        f'<xdr:cNvPicPr><a:picLocks noChangeAspect="1"/></xdr:cNvPicPr></xdr:nvPicPr>'
        f'<xdr:blipFill><a:blip xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:embed="{rid}"/>'
        f'<a:stretch><a:fillRect/></a:stretch></xdr:blipFill>'
        f'<xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr>'
        f'</xdr:pic><xdr:clientData/></xdr:oneCellAnchor>'
    )

estado_app = {}
rutas_fotos = {}

def obtener_img_state(llave):
    ruta = rutas_fotos.get(llave)
    if ruta and os.path.exists(ruta):
        return Image.open(ruta)
    return None

def generar_reporte_excel(ruta_plantilla, cliente, lugar, fecha_insp, cod_equipo,
                           cod_tolva, horometro, cod_informe, revision, pm,
                           estructura_zonas, nombre_realizado, fecha_firma, firma_ruta,
                           todos_los_rechazos, matriz_espesores=None):
    wb = openpyxl.load_workbook(ruta_plantilla)
    ws = wb["TOLVA DT"]
    
    font_black = Font(color="000000")
    font_red = Font(color="FF0000", bold=True)

    ws["C5"] = cliente
    ws["C6"] = lugar
    ws["C7"] = fecha_insp
    ws["H6"] = cod_equipo
    ws["H7"] = cod_tolva
    ws["L6"] = horometro
    ws["P5"] = cod_informe
    ws["P6"] = revision
    ws["P7"] = pm

    # --- MEDICIÓN DE ESPESORES (MATRIZ 8x7 EN COLUMNAS K a R JUNTO A ZONA 1) ---
    matriz_espesores_omitida = False
    if es_pm_espesores(pm) and matriz_espesores is not None:
        merge_fondo_diagrama = None
        for mc in ws.merged_cells.ranges:
            if mc.min_col == 1 and mc.max_col >= 18 and 8 <= mc.min_row <= 10 and (mc.max_row - mc.min_row) >= 10:
                merge_fondo_diagrama = mc
                break

        if merge_fondo_diagrama is None:
            matriz_espesores_omitida = True
        else:
            try:
                fila_ini_fondo = merge_fondo_diagrama.min_row
                fila_fin_fondo = merge_fondo_diagrama.max_row
                ws.unmerge_cells(str(merge_fondo_diagrama))
                ws.merge_cells(start_row=fila_ini_fondo, start_column=1, end_row=fila_fin_fondo, end_column=10)

                col_matriz = 11  # K
                borde = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
                borde_titulo = Border(left=Side(style="medium"), right=Side(style="medium"), top=Side(style="medium"), bottom=Side(style="medium"))
                align = Alignment(horizontal="center", vertical="center")

                alto_disponible = fila_fin_fondo - fila_ini_fondo + 1
                alto_bloque_matriz = 10
                offset_centrado = max(0, (alto_disponible - alto_bloque_matriz) // 2)
                fila_titulo_espesores = fila_ini_fondo + offset_centrado

                ws.merge_cells(start_row=fila_titulo_espesores, start_column=col_matriz, end_row=fila_titulo_espesores, end_column=col_matriz + 3)
                c_titulo = ws.cell(row=fila_titulo_espesores, column=col_matriz, value="ESPESORES DE PISO")
                c_titulo.font = Font(bold=True)
                c_titulo.alignment = align
                for col in range(col_matriz, col_matriz + 4):
                    ws.cell(row=fila_titulo_espesores, column=col).border = borde_titulo
                ws.row_dimensions[fila_titulo_espesores].height = 20

                fila_matriz = fila_titulo_espesores + 1

                ws.cell(row=fila_matriz, column=col_matriz).border = borde
                for j in range(7):
                    c = ws.cell(row=fila_matriz, column=col_matriz + 1 + j, value=j + 1)
                    c.font = Font(bold=True)
                    c.alignment = align
                    c.border = borde

                for i in range(8):
                    fila = fila_matriz + 1 + i
                    c_label = ws.cell(row=fila, column=col_matriz, value=i + 1)
                    c_label.font = Font(bold=True)
                    c_label.alignment = align
                    c_label.border = borde
                    for j in range(7):
                        valor = matriz_espesores.iloc[i, j]
                        c = ws.cell(row=fila, column=col_matriz + 1 + j, value=valor)
                        c.alignment = align
                        c.border = borde
                    ws.row_dimensions[fila].height = 18

                ws.row_dimensions[fila_matriz].height = 18

                if (matriz_espesores == "-").all().all():
                    nota = ws.cell(row=fila_matriz + 9, column=col_matriz, value="NO SE REALIZÓ MEDICIÓN DE ESPESORES")
                    ws.merge_cells(start_row=fila_matriz + 9, start_column=col_matriz, end_row=fila_matriz + 9, end_column=18)
                    nota.font = Font(bold=True, color="FF0000")
            except Exception:
                matriz_espesores_omitida = True

    desplazamiento = 0
    puntos_insercion = []
    fotos_pendientes = []

    filas_headers_zonas = _detectar_headers_zonas(ws)
    ws.row_breaks.brk = []

    presupuesto_pagina_pt = _calcular_presupuesto_pagina_pt(ws)
    altura_pagina_actual_pt = _altura_filas_pt(ws, 1, filas_headers_zonas[0] - 1)
    fila_medida_hasta = filas_headers_zonas[0] - 1
    fila_legend_zona_anterior = None
    ventanas_diagramas = []

    for idx_z, bloque_zona in enumerate(estructura_zonas):
        fila_header = filas_headers_zonas[idx_z] + desplazamiento
        fila_inicio_items = fila_header + 2
        num_items = len(bloque_zona["items"])

        if fila_legend_zona_anterior is not None:
            fila_proteccion_inicio = fila_legend_zona_anterior + 1
        else:
            fila_proteccion_inicio = fila_header - 2
        fila_legend_zona = fila_header + 2 + num_items
        altura_tabla_zona_pt = _altura_filas_pt(ws, fila_proteccion_inicio, fila_legend_zona)

        altura_pagina_actual_pt += _altura_filas_pt(ws, fila_medida_hasta + 1, fila_proteccion_inicio - 1)
        fila_medida_hasta = fila_proteccion_inicio - 1

        if altura_pagina_actual_pt + altura_tabla_zona_pt > presupuesto_pagina_pt:
            ws.row_breaks.append(Break(id=fila_proteccion_inicio - 1))
            altura_pagina_actual_pt = 0

        altura_pagina_actual_pt += altura_tabla_zona_pt
        fila_medida_hasta = fila_legend_zona
        fila_legend_zona_anterior = fila_legend_zona

        if idx_z > 0:
            ventanas_diagramas.append({
                "fila_ini": fila_proteccion_inicio,
                "fila_fin": fila_header - 2,
                "esquemas": bloque_zona["esquema"],
            })

        for idx_i, item in enumerate(bloque_zona["items"]):
            cod_z, desc_z, tec_def = item
            key_id = f"z{idx_z}_{idx_i}"
            fila = fila_inicio_items + idx_i

            defecto = estado_app.get(f"def_{key_id}", "LF")
            es_lf = (defecto == "LF")

            if es_lf:
                longitud = "-"
                est_post = "-"
                condicion = "ACEPTABLE"
                comentario = "-"
            else:
                opc_long = estado_app.get(f"opclong_{key_id}", "Manual")
                if opc_long == "VARIOS":
                    longitud = "VARIOS"
                else:
                    longitud = estado_app.get(f"longval_{key_id}", "100")
                est_post = estado_app.get(f"est_{key_id}", "NR")
                condicion = "RECHAZADO"
                comentario = estado_app.get(f"com_{key_id}", "CREAR OT")

            tecnica = estado_app.get(f"tec_{key_id}", tec_def)

            ws.cell(row=fila, column=5, value=fecha_insp)
            ws.cell(row=fila, column=7, value=defecto)
            ws.cell(row=fila, column=8, value=longitud)
            ws.cell(row=fila, column=9, value=est_post)
            ws.cell(row=fila, column=11, value=tecnica)
            ws.cell(row=fila, column=12, value=condicion)
            ws.cell(row=fila, column=14, value=comentario)

            for c_idx in [1, 2, 5, 7, 8, 9, 11, 12, 14]:
                try:
                    ws.cell(row=fila, column=c_idx).font = font_black if es_lf else font_red
                except Exception:
                    pass

        rechazos_zona = [r for r in todos_los_rechazos if r["key_id"].startswith(f"z{idx_z}_")]

        # Letreros OBLIGATORIOS de Zona 02 y Zona 08 (Siempre se generan)
        bloques_obligatorios = []
        if idx_z == 1:
            bloques_obligatorios = [
                {"descripcion": "LETRERO LATERAL RH", "llave": "esp_z2_rh"},
                {"descripcion": "LETRERO LATERAL LH", "llave": "esp_z2_lh"},
            ]
        elif idx_z == 5:
            bloques_obligatorios = [
                {"descripcion": "LETRERO POSTERIOR", "llave": "esp_z8_post"},
            ]

        if rechazos_zona or bloques_obligatorios:
            fila_legend = fila_header + 2 + num_items
            punto_insercion = fila_legend + 1

            altura_pagina_actual_pt += _altura_filas_pt(ws, fila_medida_hasta + 1, punto_insercion - 1)
            fila_medida_hasta = punto_insercion - 1

            altura_header_pt = 25
            if altura_pagina_actual_pt + altura_header_pt + 8 * 20 > presupuesto_pagina_pt:
                ws.row_breaks.append(Break(id=punto_insercion - 1))
                altura_pagina_actual_pt = 0

            _crear_header_fotos(ws, punto_insercion)
            puntos_insercion.append((punto_insercion, 1, altura_header_pt))
            desplazamiento += 1
            altura_pagina_actual_pt += altura_header_pt
            siguiente_fila_libre = punto_insercion + 1
            fila_medida_hasta = punto_insercion

            alto_bloque = 8
            altura_bloque_pt = alto_bloque * 20

            for rechazo in rechazos_zona:
                if altura_pagina_actual_pt + altura_bloque_pt > presupuesto_pagina_pt:
                    ws.row_breaks.append(Break(id=siguiente_fila_libre - 1))
                    altura_pagina_actual_pt = 0

                fila_ini, fila_fin = _crear_bloque_foto_nuevo(ws, siguiente_fila_libre, alto_bloque)
                puntos_insercion.append((siguiente_fila_libre, alto_bloque, altura_bloque_pt))
                desplazamiento += alto_bloque
                altura_pagina_actual_pt += altura_bloque_pt
                siguiente_fila_libre = fila_fin + 1
                fila_medida_hasta = fila_fin

                key_id = rechazo["key_id"]
                texto_desc = f"ZONA {rechazo['zona']}\n{rechazo['descripcion'].upper()}\n\n{_NOMBRE_COMPLETO_DEFECTO.get(rechazo['defecto'], rechazo['defecto'])}"
                ws.cell(row=fila_ini, column=1, value=texto_desc)

                for prefijo_foto, col_0idx, col_span in (("pano", 4, 7), ("det", 11, 7)):
                    llave_base = f"img_{prefijo_foto}_{key_id}"
                    img_foto = obtener_img_state(llave_base)
                    if img_foto is not None:
                        fotos_pendientes.append({
                            "fila_0idx": fila_ini - 1,
                            "col_0idx": col_0idx,
                            "col_span": col_span,
                            "row_span": alto_bloque,
                            "bytes_png": _preparar_imagen_para_insertar(img_foto),
                        })

            for oblig in bloques_obligatorios:
                if altura_pagina_actual_pt + altura_bloque_pt > presupuesto_pagina_pt:
                    ws.row_breaks.append(Break(id=siguiente_fila_libre - 1))
                    altura_pagina_actual_pt = 0

                fila_ini, fila_fin = _crear_bloque_foto_nuevo(ws, siguiente_fila_libre, alto_bloque)
                puntos_insercion.append((siguiente_fila_libre, alto_bloque, altura_bloque_pt))
                desplazamiento += alto_bloque
                altura_pagina_actual_pt += altura_bloque_pt
                siguiente_fila_libre = fila_fin + 1
                fila_medida_hasta = fila_fin

                ws.cell(row=fila_ini, column=1, value=oblig["descripcion"])
                ws.cell(row=fila_ini, column=12, value="-")  # Detalle siempre "-"

                llave_fija = oblig["llave"]
                img_pano = obtener_img_state(f"img_pano_{llave_fija}")
                if img_pano is not None:
                    fotos_pendientes.append({
                        "fila_0idx": fila_ini - 1,
                        "col_0idx": 4,
                        "col_span": 7,
                        "row_span": alto_bloque,
                        "bytes_png": _preparar_imagen_para_insertar(img_pano),
                    })

            fila_legend_zona_anterior = siguiente_fila_libre - 1

    # --- ZONAS A REPARAR (OTs) ---
    fila_ot = 148 + desplazamiento
    ws.row_breaks.append(Break(id=fila_ot - 2))
    fila_ot_max = 155 + desplazamiento
    for idx_ot, rechazo in enumerate(todos_los_rechazos):
        if fila_ot > fila_ot_max:
            break
        defecto = rechazo["defecto"]
        prefijo = "SOLD_CBO" if defecto == "DE" else "SOLD_REP"
        codigo_sugerido = f"BK00{str(idx_ot + 1).zfill(5)}"
        codigo_backlog = estado_app.get(f"bk_{rechazo['key_id']}", codigo_sugerido)
        texto_ot = f"{prefijo} {rechazo['descripcion']} ({rechazo['zona']}) - {codigo_backlog}"
        try:
            celda_ot = ws.cell(row=fila_ot, column=1, value=texto_ot)
            celda_ot.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        except Exception:
            pass
        fila_ot += 1

    fila_nombre = 157 + desplazamiento
    fila_firma = 160 + desplazamiento
    fila_fecha = 162 + desplazamiento
    try:
        ws.cell(row=fila_nombre, column=2, value=nombre_realizado)
    except Exception:
        pass
    try:
        ws.cell(row=fila_fecha, column=2, value=fecha_firma)
    except Exception:
        pass

    if firma_ruta and os.path.exists(firma_ruta):
        img_firma = Image.open(firma_ruta).convert("RGB")
        buf_firma = io.BytesIO()
        img_firma.save(buf_firma, format="PNG")
        fotos_pendientes.append({
            "fila_0idx": fila_firma - 1,
            "col_0idx": 1,
            "col_span": 3,
            "row_span": 2,
            "bytes_png": buf_firma.getvalue(),
        })

    ws.page_setup.scale = 50
    if ws.sheet_properties.pageSetUpPr is None:
        ws.sheet_properties.pageSetUpPr = PageSetupProperties()
    ws.sheet_properties.pageSetUpPr.fitToPage = False

    fila_final_original = 162
    fila_final_nueva = fila_final_original + desplazamiento
    ws.print_area = f"A1:R{fila_final_nueva}"

    buf_temp = io.BytesIO()
    wb.save(buf_temp)
    buf_temp.seek(0)
    z_temp = zipfile.ZipFile(buf_temp)
    partes_nuevas = {}
    z_original = zipfile.ZipFile(ruta_plantilla)
    
    sheet1_original = z_original.read('xl/worksheets/sheet1.xml').decode('utf-8')
    match_drawing = re.search(r'<[a-zA-Z0-9]*:?drawing r:id=".*?"\s*/>', sheet1_original)
    if not match_drawing:
        match_drawing = re.search(r'<[a-zA-Z0-9]*:?drawing r:id=".*?">.*?</[a-zA-Z0-9]*:?drawing>', sheet1_original)

    for nombre in ['xl/worksheets/sheet1.xml', 'xl/sharedStrings.xml', 'xl/styles.xml', 'xl/workbook.xml']:
        if nombre in z_temp.namelist():
            xml_data = z_temp.read(nombre).decode('utf-8')
            if nombre == 'xl/worksheets/sheet1.xml':
                xml_data = re.sub(r'<[a-zA-Z0-9]*:?drawing r:id=".*?"\s*/>', '', xml_data)
                xml_data = re.sub(r'<[a-zA-Z0-9]*:?drawing r:id=".*?">.*?</[a-zA-Z0-9]*:?drawing>', '', xml_data)
                if match_drawing:
                    tag_drawing = match_drawing.group(0)
                    if 'xmlns:r=' not in tag_drawing:
                        tag_drawing = tag_drawing.replace('r:id=', 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id=')
                    xml_data = xml_data.replace('</worksheet>', tag_drawing + '</worksheet>')
            partes_nuevas[nombre] = xml_data.encode('utf-8')

    drawing1_xml = z_original.read('xl/drawings/drawing1.xml').decode('utf-8')
    drawing1_rels = z_original.read('xl/drawings/_rels/drawing1.xml.rels').decode('utf-8')

    for texto_zona in ("ZONA 06", "ZONA 07", "ZONA 08"):
        drawing1_xml = _ajustar_altura_forma_por_texto(drawing1_xml, texto_zona, 393247)

    drawing1_xml = _estandarizar_cajas_de_puntos(drawing1_xml, 495000, 270000)

    for fila_insercion, cantidad, altura_pt in puntos_insercion:
        altura_emu = int(altura_pt * 12700)
        drawing1_xml = _desplazar_filas_drawing_xml(drawing1_xml, fila_insercion - 1, cantidad, altura_emu)

    geo = _Geometria(ws)
    rids_existentes = re.findall(r'Id="rId(\d+)"', drawing1_rels)
    siguiente_rid = max((int(r) for r in rids_existentes), default=0) + 1

    anchors_esquemas_xml = ""
    id_shape_esquema = 8000
    media_esquemas_pendientes = {}
    for ventana in ventanas_diagramas:
        fila_ini_0idx = ventana["fila_ini"] - 1
        fila_fin_0idx = ventana["fila_fin"] - 1
        if fila_fin_0idx < fila_ini_0idx:
            continue
        drawing1_xml = _quitar_anchors_en_ventana(drawing1_xml, fila_ini_0idx, fila_fin_0idx)

        nombres_png = ventana["esquemas"]
        n_imgs = len(nombres_png)
        ancho_total_cols = 18
        for idx_img, nombre_png in enumerate(nombres_png):
            ruta_png = buscar_recurso("imagenes_esquemas", nombre_png) or os.path.join("imagenes_esquemas", nombre_png)
            if not os.path.exists(ruta_png):
                continue
            col_ini = int(idx_img * ancho_total_cols / n_imgs)
            col_fin = int((idx_img + 1) * ancho_total_cols / n_imgs)
            with open(ruta_png, "rb") as f_img:
                datos_png = f_img.read()
            nombre_media = f"esquema{siguiente_rid}.png"
            rid_actual = f"rId{siguiente_rid}"
            drawing1_rels = re.sub(
                r'</([a-zA-Z0-9]+:)?Relationships>',
                lambda m, nm=nombre_media, rid=rid_actual: f'<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/{nm}" Id="{rid}"/>' + m.group(0),
                drawing1_rels
            )
            anchors_esquemas_xml += _anchor_imagen_ajustada_xml(
                geo, id_shape_esquema, "EsquemaZona", rid_actual, datos_png,
                col_ini, col_fin, fila_ini_0idx, fila_fin_0idx + 1, margen=40000
            )
            media_esquemas_pendientes[f"xl/media/{nombre_media}"] = datos_png
            siguiente_rid += 1
            id_shape_esquema += 1

    drawing1_xml = re.sub(
        r'</([a-zA-Z0-9]+:)?wsDr>',
        lambda m: anchors_esquemas_xml + m.group(0),
        drawing1_xml
    )

    media_nuevos = dict(media_esquemas_pendientes)
    anchors_nuevos_xml = ""
    id_shape = 9000

    for foto in fotos_pendientes:
        nombre_media = f"appFoto{siguiente_rid}.png"
        media_nuevos[f"xl/media/{nombre_media}"] = foto["bytes_png"]
        rid_actual = f"rId{siguiente_rid}"
        
        drawing1_rels = re.sub(
            r'</([a-zA-Z0-9]+:)?Relationships>',
            lambda m, nm=nombre_media, rid=rid_actual: f'<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/{nm}" Id="{rid}"/>' + m.group(0),
            drawing1_rels
        )
        
        anchors_nuevos_xml += _anchor_imagen_ajustada_xml(
            geo, id_shape, "FotoApp", rid_actual, foto["bytes_png"],
            foto["col_0idx"], foto["col_0idx"] + foto["col_span"],
            foto["fila_0idx"], foto["fila_0idx"] + foto["row_span"], margen=100000
        )
        siguiente_rid += 1
        id_shape += 1

    drawing1_xml = re.sub(
        r'</([a-zA-Z0-9]+:)?wsDr>',
        lambda m: anchors_nuevos_xml + m.group(0),
        drawing1_xml
    )
    
    partes_nuevas['xl/drawings/drawing1.xml'] = drawing1_xml.encode('utf-8')
    partes_nuevas['xl/drawings/_rels/drawing1.xml.rels'] = drawing1_rels.encode('utf-8')

    buf_final = io.BytesIO()
    with zipfile.ZipFile(buf_final, 'w', zipfile.ZIP_DEFLATED) as z_final:
        for item in z_original.infolist():
            if item.filename == 'xl/calcChain.xml':
                continue
            datos = partes_nuevas.get(item.filename, z_original.read(item.filename))
            z_final.writestr(item, datos)
        for nombre_media, datos_media in media_nuevos.items():
            z_final.writestr(nombre_media, datos_media)
    buf_final.seek(0)
    avisos = []
    if matriz_espesores_omitida:
        avisos.append(
            "⚠️ No se pudo colocar la matriz de espesores en el Excel: no se encontró "
            "el área combinada del diagrama de Zona 1 en la plantilla."
        )
    return buf_final.getvalue(), avisos


# ==============================================================================
# 3. INTERFAZ GRÁFICA EN FLET 1.0+ (100% FIEL A STREAMLIT)
# ==============================================================================
def main(page: ft.Page):
    page.title = "Sistema de Inspección de Tolvas CAT 794 AC"
    page.scroll = "adaptive"
    page.padding = 12
    page.theme_mode = ft.ThemeMode.DARK
    page.bgcolor = ft.Colors.with_opacity(1, "#0e1117")

    db_tolvas = cargar_db()
    controles = {}
    llave_foto_actual = None

    # --- SELECTOR DE ARCHIVOS Y CÁMARA (FLET 1.0, ASÍNCRONO) ---
    selector_fotos = ft.FilePicker()
    try:
        page.services.append(selector_fotos)
    except Exception:
        page.overlay.append(selector_fotos)

    def es_movil():
        try:
            return bool(page.platform.is_mobile())
        except Exception:
            return page.platform in (ft.PagePlatform.ANDROID, ft.PagePlatform.IOS)

    def marcar_foto(llave, ruta, nombre):
        rutas_fotos[llave] = ruta
        if llave in controles:
            controles[llave].value = f"✅ {nombre}"
            controles[llave].color = ft.Colors.GREEN_400
        page.update()

    def marcar_error(llave, texto):
        if llave in controles:
            controles[llave].value = f"❌ {texto}"
            controles[llave].color = ft.Colors.RED_400
        page.update()

    MSG_FORMATO = "Formato de imagen no compatible (¿HEIC?). Usa 📷 Cámara, o en la cámara del teléfono elige formato JPEG / 'más compatible'"

    async def seleccionar_async(llave):
        try:
            archivos = await selector_fotos.pick_files(
                allow_multiple=False,
                file_type=ft.FilePickerFileType.IMAGE,
            )
            if not archivos:
                return
            f0 = archivos[0]
            if not f0.path:
                marcar_error(llave, "No se pudo leer la ruta del archivo")
                return
            with open(f0.path, "rb") as fh:
                datos = fh.read()
            if llave == "firma_upload":
                try:
                    abrir_imagen(datos)
                except Exception:
                    marcar_error(llave, MSG_FORMATO)
                    return
                marcar_foto(llave, f0.path, f0.name)
            else:
                procesar_foto(llave, datos, f0.name)
        except Exception as ex:
            marcar_error(llave, f"Selector: {ex}")

    def abrir_selector(llave):
        page.run_task(seleccionar_async, llave)

    # --- Cambios de tamaño / rotación de pantalla ---
    resize_cbs = []

    def _on_resize(e=None):
        for cb in list(resize_cbs):
            try:
                cb()
            except Exception:
                pass

    try:
        page.on_resize = _on_resize
    except Exception:
        pass

    def tam_pantalla():
        return (page.width or 360), (page.height or 640)

    # --- Utilidades de imagen ---
    def abrir_imagen(datos):
        im = Image.open(io.BytesIO(datos))
        return ImageOps.exif_transpose(im).convert("RGB")

    def guardar_jpeg(llave, im):
        carpeta = os.environ.get("FLET_APP_STORAGE_TEMP") or tempfile.gettempdir()
        os.makedirs(carpeta, exist_ok=True)
        ruta = os.path.join(carpeta, f"{llave}.jpg")
        im.save(ruta, "JPEG", quality=88)
        return ruta

    def procesar_foto(llave, datos, nombre):
        try:
            im = abrir_imagen(datos)
            im.thumbnail((2000, 2000))
        except Exception:
            marcar_error(llave, MSG_FORMATO)
            return
        abrir_editor(llave, im, nombre)

    # --- Editor de anotaciones (rectángulo, círculo, flecha, lápiz) ---
    def abrir_editor(llave, im, nombre):
        iw, ih = im.size
        vista = im.copy()
        vista.thumbnail((1400, 1400))
        buf = io.BytesIO()
        vista.save(buf, "JPEG", quality=80)
        grosor_img = max(4, iw // 200)
        COLORES = ["#FF1744", "#FFEA00", "#00E676", "#FFFFFF"]
        est = {"herr": "rect", "color": COLORES[0], "actual": None, "dw": 1.0, "dh": 1.0}
        formas = []

        img_ctl = ft.Image(src=buf.getvalue(), fit=ft.BoxFit.FILL)

        def clamp01(v):
            return max(0.0, min(1.0, v))

        def punta_flecha(x0, y0, x1, y1, largo):
            ang = math.atan2(y1 - y0, x1 - x0)
            t = math.radians(25)
            p1 = (x1 - largo * math.cos(ang - t), y1 - largo * math.sin(ang - t))
            p2 = (x1 - largo * math.cos(ang + t), y1 - largo * math.sin(ang + t))
            return p1, p2

        def a_cv(f):
            dw, dh = est["dw"], est["dh"]
            g = max(2.5, grosor_img * dw / iw)
            paint = ft.Paint(color=f["color"], stroke_width=g, style=ft.PaintingStyle.STROKE)
            d = f["datos"]
            if f["tipo"] == "lapiz":
                pts = [(x * dw, y * dh) for x, y in d]
                if len(pts) < 2:
                    return []
                els = [cv.Path.MoveTo(pts[0][0], pts[0][1])] + [cv.Path.LineTo(x, y) for x, y in pts[1:]]
                return [cv.Path(els, paint=paint)]
            x0, y0, x1, y1 = d[0] * dw, d[1] * dh, d[2] * dw, d[3] * dh
            if f["tipo"] == "rect":
                return [cv.Rect(min(x0, x1), min(y0, y1), abs(x1 - x0), abs(y1 - y0), paint=paint)]
            if f["tipo"] == "oval":
                return [cv.Oval(min(x0, x1), min(y0, y1), abs(x1 - x0), abs(y1 - y0), paint=paint)]
            p1, p2 = punta_flecha(x0, y0, x1, y1, max(14.0, g * 5))
            return [
                cv.Line(x0, y0, x1, y1, paint=paint),
                cv.Line(x1, y1, p1[0], p1[1], paint=paint),
                cv.Line(x1, y1, p2[0], p2[1], paint=paint),
            ]

        def dibujar():
            shapes = []
            for f in formas + ([est["actual"]] if est["actual"] else []):
                shapes.extend(a_cv(f))
            lienzo.shapes = shapes
            lienzo.update()

        def pan_start(e):
            x = clamp01(e.local_position.x / est["dw"])
            y = clamp01(e.local_position.y / est["dh"])
            if est["herr"] == "lapiz":
                est["actual"] = {"tipo": "lapiz", "color": est["color"], "datos": [(x, y)]}
            else:
                est["actual"] = {"tipo": est["herr"], "color": est["color"], "datos": [x, y, x, y]}

        def pan_update(e):
            f = est["actual"]
            if not f:
                return
            x = clamp01(e.local_position.x / est["dw"])
            y = clamp01(e.local_position.y / est["dh"])
            if f["tipo"] == "lapiz":
                f["datos"].append((x, y))
            else:
                f["datos"][2], f["datos"][3] = x, y
            dibujar()

        def pan_end(e=None):
            f = est["actual"]
            est["actual"] = None
            if f:
                d = f["datos"]
                valido = len(d) > 3 if f["tipo"] == "lapiz" else (abs(d[2] - d[0]) > 0.006 or abs(d[3] - d[1]) > 0.006)
                if valido:
                    formas.append(f)
            dibujar()

        lienzo = cv.Canvas(
            shapes=[],
            content=ft.GestureDetector(
                on_pan_start=pan_start, on_pan_update=pan_update, on_pan_end=pan_end, drag_interval=15
            ),
        )
        area = ft.Stack([img_ctl, lienzo])

        # --- barra de herramientas ---
        botones_herr = {}

        def elegir_herr(h):
            est["herr"] = h
            for k, b_ in botones_herr.items():
                b_.icon_color = ft.Colors.AMBER_400 if k == h else ft.Colors.WHITE
            page.update()

        for clave, icono, tip in [
            ("rect", ft.Icons.CROP_SQUARE, "Rectángulo"),
            ("oval", ft.Icons.PANORAMA_FISH_EYE, "Círculo"),
            ("flecha", ft.Icons.NORTH_EAST, "Flecha"),
            ("lapiz", ft.Icons.EDIT, "Lápiz"),
        ]:
            botones_herr[clave] = ft.IconButton(
                icon=icono, tooltip=tip, icon_color=ft.Colors.WHITE,
                on_click=lambda e, c=clave: elegir_herr(c),
            )
        botones_herr["rect"].icon_color = ft.Colors.AMBER_400

        chips = {}

        def elegir_color(c):
            est["color"] = c
            for k, ch in chips.items():
                ch.border = ft.Border.all(3, ft.Colors.WHITE if k == c else ft.Colors.TRANSPARENT)
            page.update()

        for c in COLORES:
            chips[c] = ft.Container(
                width=30, height=30, bgcolor=c, border_radius=15,
                border=ft.Border.all(3, ft.Colors.WHITE if c == COLORES[0] else ft.Colors.TRANSPARENT),
                on_click=lambda e, c=c: elegir_color(c),
            )

        def deshacer(e=None):
            if formas:
                formas.pop()
                dibujar()

        def borrar_todo(e=None):
            formas.clear()
            dibujar()

        def cerrar_editor(e=None):
            if raiz in page.overlay:
                page.overlay.remove(raiz)
            if ubicar in resize_cbs:
                resize_cbs.remove(ubicar)
            page.update()

        def guardar(anotar):
            try:
                salida = im.copy()
                if anotar and formas:
                    d = ImageDraw.Draw(salida)
                    for f in formas:
                        col = f["color"]
                        dd = f["datos"]
                        if f["tipo"] == "lapiz":
                            pts = [(x * iw, y * ih) for x, y in dd]
                            if len(pts) > 1:
                                d.line(pts, fill=col, width=grosor_img, joint="curve")
                            continue
                        x0, y0, x1, y1 = dd[0] * iw, dd[1] * ih, dd[2] * iw, dd[3] * ih
                        if f["tipo"] == "rect":
                            d.rectangle([min(x0, x1), min(y0, y1), max(x0, x1), max(y0, y1)], outline=col, width=grosor_img)
                        elif f["tipo"] == "oval":
                            d.ellipse([min(x0, x1), min(y0, y1), max(x0, x1), max(y0, y1)], outline=col, width=grosor_img)
                        else:
                            p1, p2 = punta_flecha(x0, y0, x1, y1, max(14.0, grosor_img * 5))
                            d.line([(x0, y0), (x1, y1)], fill=col, width=grosor_img)
                            d.line([(x1, y1), p1], fill=col, width=grosor_img)
                            d.line([(x1, y1), p2], fill=col, width=grosor_img)
                ruta = guardar_jpeg(llave, salida)
                cerrar_editor()
                marcar_foto(llave, ruta, nombre + (" (anotada)" if anotar and formas else ""))
            except Exception as ex:
                marcar_error(llave, f"No se pudo guardar la foto: {ex}")
                cerrar_editor()

        barra = ft.Container(
            bgcolor=ft.Colors.with_opacity(0.85, ft.Colors.BLACK), padding=8,
            content=ft.Column([
                ft.Row(list(botones_herr.values()) + [
                    ft.IconButton(icon=ft.Icons.UNDO, tooltip="Deshacer", icon_color=ft.Colors.WHITE, on_click=deshacer),
                    ft.IconButton(icon=ft.Icons.DELETE_OUTLINE, tooltip="Borrar todo", icon_color=ft.Colors.WHITE, on_click=borrar_todo),
                ], wrap=True, spacing=0, run_spacing=0),
                ft.Row(list(chips.values()), spacing=10),
                ft.Row([
                    ft.FilledButton("Guardar", icon=ft.Icons.CHECK, on_click=lambda e: guardar(True)),
                    ft.OutlinedButton("Sin anotar", on_click=lambda e: guardar(False)),
                    ft.TextButton("Cancelar", on_click=cerrar_editor),
                ], wrap=True, spacing=6, run_spacing=4),
            ], spacing=6),
        )
        caja = ft.Container(content=area, alignment=ft.Alignment.CENTER, left=0, top=0, right=0, bottom=0)

        def ubicar():
            W, H = tam_pantalla()
            horiz = W > H
            if horiz:
                disp_w, disp_h = W - 240, H - 24
                barra.left, barra.right, barra.top, barra.bottom, barra.width = None, 0, 0, 0, 230
                caja.right, caja.bottom = 230, 0
            else:
                disp_w, disp_h = W - 16, H - 230
                barra.left, barra.right, barra.top, barra.bottom, barra.width = 0, 0, None, 0, None
                caja.right, caja.bottom = 0, 200
            esc = max(0.05, min(disp_w / iw, disp_h / ih))
            dw, dh = iw * esc, ih * esc
            est["dw"], est["dh"] = dw, dh
            for c_ in (img_ctl, lienzo, area):
                c_.width, c_.height = dw, dh
            dibujar()

        raiz = ft.Container(
            expand=True, bgcolor=ft.Colors.BLACK,
            content=ft.SafeArea(content=ft.Stack([caja, barra], expand=True), expand=True),
        )
        resize_cbs.append(ubicar)
        page.overlay.append(raiz)
        ubicar()
        page.update()

    # --- Cámara a pantalla completa con zoom y selector de lente ---
    async def camara_async(llave):
        try:
            import flet_camera as fc
        except Exception:
            marcar_error(llave, "Cámara no disponible en esta versión; usa Galería")
            return

        camara = fc.Camera(expand=True, preview_enabled=True)
        estado_cam = ft.Text("Iniciando cámara...", size=12, color=ft.Colors.WHITE)
        btn_tomar = ft.FilledButton("Tomar foto", icon=ft.Icons.PHOTO_CAMERA, disabled=True)
        etq_zoom = ft.Text("Zoom 1.0x", size=12, color=ft.Colors.WHITE)
        sld_zoom = ft.Slider(min=1, max=2, value=1, disabled=True)
        fila_zoom = ft.Row(wrap=True, spacing=4, run_spacing=4)
        dd_cam = ft.Dropdown(label="Cámara / lente", width=210, dense=True, text_size=13, visible=False)
        lista_camaras = []

        panel = ft.Container(
            bgcolor=ft.Colors.with_opacity(0.65, ft.Colors.BLACK), padding=10,
            content=ft.Column([
                estado_cam, etq_zoom, sld_zoom, fila_zoom, dd_cam,
                ft.Row([btn_tomar, ft.TextButton("Cancelar", on_click=lambda e: cerrar())], wrap=True, spacing=8),
            ], spacing=6, scroll=ft.ScrollMode.AUTO),
        )
        fondo = ft.Container(content=camara, bgcolor=ft.Colors.BLACK, left=0, top=0, right=0, bottom=0)

        def ubicar():
            W, H = tam_pantalla()
            if W > H:
                panel.left, panel.right, panel.top, panel.bottom, panel.width = None, 0, 0, 0, 230
            else:
                panel.left, panel.right, panel.top, panel.bottom, panel.width = 0, 0, None, 0, None

        raiz = ft.Container(
            expand=True, bgcolor=ft.Colors.BLACK,
            content=ft.SafeArea(content=ft.Stack([fondo, panel], expand=True), expand=True),
        )

        def cerrar():
            if raiz in page.overlay:
                page.overlay.remove(raiz)
            if _redib in resize_cbs:
                resize_cbs.remove(_redib)
            page.update()

        def _redib():
            ubicar()
            page.update()

        async def poner_zoom(z):
            try:
                sld_zoom.value = z
                etq_zoom.value = f"Zoom {z:.1f}x"
                page.update()
                await camara.set_zoom_level(z)
            except Exception as ex:
                estado_cam.value = f"Zoom no disponible: {ex}"
                page.update()

        async def cambiar_zoom(e):
            await poner_zoom(float(e.control.value))

        sld_zoom.on_change = cambiar_zoom

        def etiqueta_camara(c, n):
            if c.lens_direction == fc.CameraLensDirection.BACK:
                base = "Trasera"
            elif c.lens_direction == fc.CameraLensDirection.FRONT:
                base = "Frontal"
            else:
                base = "Externa"
            extra = " (gran angular)" if getattr(getattr(c, "lens_type", None), "value", "") == "ultraWide" else ""
            return f"{base} {n}{extra}"

        async def iniciar(desc):
            btn_tomar.disabled = True
            estado_cam.value = "Iniciando cámara..."
            page.update()
            await camara.initialize(
                description=desc,
                resolution_preset=fc.ResolutionPreset.HIGH,
                enable_audio=False,
                image_format_group=fc.ImageFormatGroup.JPEG,
            )
            try:
                mn = float(await camara.get_min_zoom_level())
                mx = min(float(await camara.get_max_zoom_level()), 10.0)
            except Exception:
                mn, mx = 1.0, 1.0
            if mx - mn > 0.05:
                sld_zoom.min, sld_zoom.max = mn, mx
                sld_zoom.value = min(max(1.0, mn), mx)
                sld_zoom.disabled = False
            else:
                sld_zoom.min, sld_zoom.max, sld_zoom.value = 1, 2, 1
                sld_zoom.disabled = True
            etq_zoom.value = f"Zoom {float(sld_zoom.value):.1f}x"
            fila_zoom.controls = [
                ft.OutlinedButton(f"{z:g}x", on_click=lambda e, z=z: page.run_task(poner_zoom, float(z)))
                for z in (0.5, 1, 2, 3, 5) if mn - 0.01 <= z <= mx + 0.01 and mx - mn > 0.05
            ]
            btn_tomar.disabled = False
            estado_cam.value = "Cámara lista"
            page.update()

        async def tomar(e=None):
            try:
                btn_tomar.disabled = True
                estado_cam.value = "Capturando..."
                page.update()
                datos = await camara.take_picture()
                cerrar()
                procesar_foto(llave, datos, "foto de cámara")
            except Exception as ex:
                estado_cam.value = f"Error al capturar: {ex}"
                btn_tomar.disabled = False
                page.update()

        async def cambiar_camara(e):
            c = next((x for x in lista_camaras if x.name == dd_cam.value), None)
            if c:
                try:
                    await iniciar(c)
                except Exception as ex:
                    estado_cam.value = f"No se pudo cambiar de cámara: {ex}"
                    page.update()

        btn_tomar.on_click = tomar
        dd_cam.on_select = cambiar_camara
        resize_cbs.append(_redib)
        ubicar()
        page.overlay.append(raiz)
        page.update()

        try:
            lista_camaras.extend(await camara.get_available_cameras())
            if not lista_camaras:
                estado_cam.value = "No se encontró ninguna cámara"
                page.update()
                return
            elegida = next((c for c in lista_camaras if c.lens_direction == fc.CameraLensDirection.BACK), lista_camaras[0])
            if len(lista_camaras) > 1:
                dd_cam.options = [ft.DropdownOption(key=c.name, text=etiqueta_camara(c, i + 1)) for i, c in enumerate(lista_camaras)]
                dd_cam.value = elegida.name
                dd_cam.visible = True
            await iniciar(elegida)
        except Exception as ex:
            estado_cam.value = f"No se pudo iniciar la cámara: {ex}"
            page.update()

    def abrir_camara(llave):
        page.run_task(camara_async, llave)

    def boton_foto(titulo, llave):
        return ft.Column([
            ft.Text(titulo, size=12, weight="bold"),
            ft.Row([
                ft.Button("📷 Cámara", on_click=lambda e, k=llave: abrir_camara(k)),
                ft.OutlinedButton("🖼️ Galería", on_click=lambda e, k=llave: abrir_selector(k)),
            ], wrap=True, spacing=8, run_spacing=6),
        ], spacing=4)

    # --- TÍTULO E INFO ---
    page.add(
        ft.Text("📋 Reporte de Inspección de Tolvas CAT 794 AC", size=22, weight="bold", color=ft.Colors.BLUE_200),
        ft.Container(
            content=ft.Text(
                "💡 IMPORTANTE: Asegúrate de que tu archivo plantilla_tolva.xlsx esté completamente vacío de fotos. "
                "Si la plantilla tiene fotos de ejemplo, bórralas, de lo contrario se mezclarán con las que selecciones aquí.",
                color=ft.Colors.WHITE, size=13
            ),
            bgcolor=ft.Colors.BLUE_900, padding=12, border_radius=8,
            margin=ft.Margin(left=0, top=5, right=0, bottom=20)
        ),
        ft.Text("1. Datos Generales del Informe", size=22, weight="bold", color=ft.Colors.WHITE)
    )

    # --- SECCIÓN 1: ENCABEZADO ---
    hoy_str = date.today().strftime("%d/%m/%Y")

    # Columna 1: Cliente, Lugar, Fecha
    dd_cliente = ft.Dropdown(
        label="Cliente:",
        options=[ft.dropdown.Option("ANGLOAMERICAN QUELLAVECO S.A."), ft.dropdown.Option("[Entrada Manual]")],
        value="ANGLOAMERICAN QUELLAVECO S.A.", dense=True
    )
    txt_cliente_manual = ft.TextField(label="Nombre del Cliente:", value="ANGLOAMERICAN QUELLAVECO S.A.", visible=False, dense=True)

    def al_cambiar_cliente(e):
        txt_cliente_manual.visible = (dd_cliente.value == "[Entrada Manual]")
        page.update()
    dd_cliente.on_select = al_cambiar_cliente

    dd_lugar = ft.Dropdown(
        label="Lugar:",
        options=[ft.dropdown.Option("TRUCK SHOP"), ft.dropdown.Option("[Entrada Manual]")],
        value="TRUCK SHOP", dense=True
    )
    txt_lugar_manual = ft.TextField(label="Lugar de Inspección:", value="TRUCK SHOP", visible=False, dense=True)

    def al_cambiar_lugar(e):
        txt_lugar_manual.visible = (dd_lugar.value == "[Entrada Manual]")
        page.update()
    dd_lugar.on_select = al_cambiar_lugar

    txt_fecha_insp = ft.TextField(label="Fecha de Inspección (DD/MM/AAAA):", value=hoy_str, dense=True)

    # Columna 2: Equipo, Tolva (con memoria JSON), Horómetro
    lista_equipos = [f"HT{str(i).zfill(3)}" for i in range(1, 35)]
    dd_equipo = ft.Dropdown(
        label="Código de Equipo:",
        options=[ft.dropdown.Option(eq) for eq in lista_equipos],
        value="HT002", dense=True
    )
    tolva_ini = db_tolvas.get("HT002", "T-CA2")
    txt_tolva = ft.TextField(label="Código de Tolva:", value=tolva_ini, dense=True)

    def guardar_tolva_db(e=None):
        eq = dd_equipo.value
        val = txt_tolva.value.strip()
        if val and val != db_tolvas.get(eq):
            db_tolvas[eq] = val
            guardar_db(db_tolvas)

    txt_tolva.on_change = guardar_tolva_db
    txt_horometro = ft.TextField(label="Horómetro (hrs):", value="34876.3", dense=True)

    # Columna 3: Sufijo, Código Informe, Revisión, PM (con opción manual y múltiplos de 1000)
    txt_sufijo = ft.TextField(label="Sufijo del Informe (XXX):", value="023", dense=True)
    lbl_cod_informe = ft.Text(f"Código Generado: ZA-IF-HT002-023", weight="bold", color=ft.Colors.BLUE_300, size=13)

    def actualizar_codigo_informe(e=None):
        eq = dd_equipo.value
        suf = txt_sufijo.value.strip()
        lbl_cod_informe.value = f"Código Generado: ZA-IF-{eq}-{suf}"
        if e and e.control == dd_equipo:
            txt_tolva.value = db_tolvas.get(eq, "T-CA2" if eq == "HT002" else "")
        page.update()

    dd_equipo.on_select = actualizar_codigo_informe
    txt_sufijo.on_change = actualizar_codigo_informe

    txt_revision = ft.TextField(label="Revisión:", value="00", dense=True)

    dd_pm = ft.Dropdown(
        label="Mantenimiento (PM):",
        options=[ft.dropdown.Option(x) for x in ["500H", "1000H", "1500H", "2000H", "3000H", "4000H", "[Entrada Manual]"]],
        value="500H", dense=True
    )
    txt_pm_manual = ft.TextField(label="PM Manual (ej. 1000, 2000, 3000):", value="1000", visible=False, dense=True)

    page.add(
        ft.ResponsiveRow([
            ft.Column([dd_cliente, txt_cliente_manual, dd_lugar, txt_lugar_manual, txt_fecha_insp], col={"sm": 12, "md": 4}),
            ft.Column([dd_equipo, txt_tolva, txt_horometro], col={"sm": 12, "md": 4}),
            ft.Column([
                txt_sufijo,
                ft.Container(content=lbl_cod_informe, bgcolor=ft.Colors.BLUE_GREY_900, padding=8, border_radius=6),
                txt_revision,
                dd_pm,
                txt_pm_manual
            ], col={"sm": 12, "md": 4}),
        ]),
        ft.Divider(height=25, color=ft.Colors.BLUE_GREY_700)
    )

    # ==========================================================================
    # SECCIÓN MEDICIÓN DE ESPESORES: MATRIZ 8x7 REAL (Punto 1..8 x Eje 1..7)
    # ==========================================================================
    chk_no_medicion = ft.Checkbox(
        label="⚠️ Marcar como 'NO SE REALIZÓ MEDICIÓN DE ESPESORES'",
        value=False
    )
    aviso_no_medicion = ft.Container(
        content=ft.Text(
            "⚠️ Se ha seleccionado omitir la medición. El reporte se llenará con '-' y agregará la nota explicativa.",
            color=ft.Colors.AMBER_300, size=13
        ),
        bgcolor=ft.Colors.with_opacity(0.2, ft.Colors.AMBER_900),
        padding=10, border_radius=6, visible=False
    )

    celdas_matriz = []
    filas_ui_matriz = []

    # Encabezado de la tabla 8x7 (Esquina + Eje 1 a Eje 7)
    header_cols = [ft.Container(content=ft.Text("Punto / Eje", weight="bold", size=12, color=ft.Colors.BLUE_200), width=90)]
    for j in range(7):
        header_cols.append(
            ft.Container(
                content=ft.Text(f"Eje {j+1}", weight="bold", size=12, color=ft.Colors.BLUE_200),
                width=85, alignment=ft.Alignment(0, 0)
            )
        )
    filas_ui_matriz.append(ft.Row(header_cols, spacing=6))

    # 8 Filas (Punto 1 a Punto 8) x 7 Columnas
    for i in range(8):
        fila_celdas = []
        controles_fila = [
            ft.Container(
                content=ft.Text(f"Punto {i+1}", weight="bold", size=12, color=ft.Colors.GREY_300),
                width=90, alignment=ft.Alignment(-1, 0)
            )
        ]
        for j in range(7):
            tf = ft.TextField(
                value="20.00",
                width=85,
                dense=True,
                text_size=12,
                text_align=ft.TextAlign.CENTER
            )
            fila_celdas.append(tf)
            controles_fila.append(tf)
        celdas_matriz.append(fila_celdas)
        filas_ui_matriz.append(ft.Row(controles_fila, spacing=6))

    def al_cambiar_no_medicion(e):
        omitir = chk_no_medicion.value
        aviso_no_medicion.visible = omitir
        for i in range(8):
            for j in range(7):
                celdas_matriz[i][j].disabled = omitir
                celdas_matriz[i][j].value = "-" if omitir else "20.00"
        page.update()

    chk_no_medicion.on_change = al_cambiar_no_medicion

    contenedor_matriz_8x7 = ft.Container(
        content=ft.Column([
            ft.Text("📏 Mapeo / Medición de Espesores (Matriz 8x7)", size=22, weight="bold", color=ft.Colors.BLUE_200),
            chk_no_medicion,
            aviso_no_medicion,
            ft.Text("Ingrese manualmente las lecturas de ultrasonido (mm) en la matriz:", size=12, color=ft.Colors.GREY_400),
            ft.Container(
                content=ft.Row([ft.Column(filas_ui_matriz, spacing=6)], scroll=ft.ScrollMode.AUTO),
                padding=12, bgcolor=ft.Colors.with_opacity(1, "#161a24"), border_radius=8
            ),
            ft.Divider(height=25, color=ft.Colors.BLUE_GREY_700)
        ], spacing=10),
        visible=False  # Inicia oculto porque el PM inicial es 500H
    )

    page.add(contenedor_matriz_8x7)

    def obtener_pm_efectivo():
        return txt_pm_manual.value.strip() if dd_pm.value == "[Entrada Manual]" else dd_pm.value

    def evaluar_visibilidad_matriz(e=None):
        txt_pm_manual.visible = (dd_pm.value == "[Entrada Manual]")
        pm_actual = obtener_pm_efectivo()
        contenedor_matriz_8x7.visible = es_pm_espesores(pm_actual)
        page.update()

    dd_pm.on_select = evaluar_visibilidad_matriz
    txt_pm_manual.on_change = evaluar_visibilidad_matriz

    # --- ESQUEMA VISUAL GENERAL ---
    for ext in ("esquema_tolva.png", "esquema_tolva.jpg"):
        ruta_ext = buscar_recurso(ext)
        if ruta_ext:
            with open(ruta_ext, "rb") as f_ext:
                bytes_ext = f_ext.read()
            page.add(
                ft.Text("🗺️ Esquema Guía General de Zonas", size=18, weight="bold", color=ft.Colors.WHITE),
                ft.Image(src=bytes_ext, height=280, fit="contain", border_radius=8),
                ft.Text("Plano de Ubicación General de Componentes - Tolva CAT 794 AC", size=11, color=ft.Colors.GREY_400),
                ft.Divider(height=20, color=ft.Colors.BLUE_GREY_800)
            )
            break

    # ==========================================================================
    # SECCIÓN 3: DESPLIEGUE DE LAS 8 ZONAS Y SUS 9 COLUMNAS FIELES
    # ==========================================================================
    contenedor_resumen_ots = ft.Column(spacing=8)

    def reconstruir_resumen_ots():
        contenedor_resumen_ots.controls.clear()
        rechazos_actuales = []
        for idx_z, bz in enumerate(ESTRUCTURA_ZONAS):
            for idx_i, it in enumerate(bz["items"]):
                cod_z, desc_z, _ = it
                k_id = f"z{idx_z}_{idx_i}"
                val_def = controles[f"def_{k_id}"].value
                if val_def != "LF":
                    rechazos_actuales.append((cod_z, desc_z, val_def, k_id))

        if not rechazos_actuales:
            contenedor_resumen_ots.controls.append(
                ft.Container(
                    content=ft.Text("✔ No se registraron defectos. No hay reparaciones pendientes requeridas.", color=ft.Colors.GREEN_300, weight="bold"),
                    bgcolor=ft.Colors.with_opacity(0.15, ft.Colors.GREEN_900), padding=12, border_radius=6
                )
            )
        else:
            contenedor_resumen_ots.controls.append(
                ft.Text("Resumen sugerido de Órdenes de Trabajo según los defectos reportados:", size=12, color=ft.Colors.GREY_400)
            )
            for idx_ot, (cod_z, desc_z, val_def, k_id) in enumerate(rechazos_actuales):
                prefijo = "SOLD_CBO" if val_def == "DE" else "SOLD_REP"
                sugerido = f"BK00{str(idx_ot + 1).zfill(5)}"
                llave_bk = f"bk_{k_id}"
                if llave_bk not in controles:
                    controles[llave_bk] = ft.TextField(
                        label=f"Código Backlog SAP ({cod_z}):",
                        value=sugerido, dense=True, width=220, text_size=13
                    )
                contenedor_resumen_ots.controls.append(
                    ft.Column([
                        ft.Container(
                            content=ft.Text(f"{prefijo} {desc_z} ({cod_z}) - {val_def}", font_family="Consolas", size=13, color=ft.Colors.BLUE_200),
                            bgcolor=ft.Colors.BLUE_GREY_900, padding=10, border_radius=6
                        ),
                        controles[llave_bk]
                    ], spacing=6)
                )

    for idx_z, zona in enumerate(ESTRUCTURA_ZONAS):
        page.add(ft.Text(zona["titulo"], size=22, weight="bold", color=ft.Colors.BLUE_200))
        
        imagenes_row = ft.Row(spacing=20, wrap=True)
        for img_name in zona["esquema"]:
            ruta_local = buscar_recurso("imagenes_esquemas", img_name)
            if ruta_local:
                with open(ruta_local, "rb") as f_loc:
                    bytes_loc = f_loc.read()
                imagenes_row.controls.append(
                    ft.Image(src=bytes_loc, height=220, fit="contain", border_radius=8)
                )
            else:
                imagenes_row.controls.append(
                    ft.Text(f"⚠️ No se encontró el esquema '{img_name}' (carpeta imagenes_esquemas)", size=12, color=ft.Colors.AMBER_400)
                )
        if imagenes_row.controls:
            page.add(
                ft.Text("🗺️ Esquema de referencia de la zona:", size=13, weight="bold", color=ft.Colors.GREY_300),
                imagenes_row,
                ft.Divider(height=10, color=ft.Colors.TRANSPARENT)
            )
        
        for idx_i, item in enumerate(zona["items"]):
            cod_z, desc_z, tec_def = item
            key_id = f"z{idx_z}_{idx_i}"
            
            lbl_cond = ft.Text("ACEPTABLE", weight="bold", size=13, color=ft.Colors.GREEN_400)
            
            dd_opclong = ft.Dropdown(
                options=[ft.dropdown.Option("Manual"), ft.dropdown.Option("VARIOS")],
                value="Manual", dense=True, text_size=13, visible=False, width=150, label="Tipo de longitud"
            )
            txt_long = ft.TextField(value="-", disabled=True, dense=True, text_size=13, width=150, label="Longitud (mm)")
            col_longitud = ft.Column([dd_opclong, txt_long], spacing=4, width=150)

            def crear_evento_opclong(dd_ol, txt_l):
                def al_cambiar_ol(e):
                    if dd_ol.value == "VARIOS":
                        txt_l.visible = False
                    else:
                        txt_l.visible = True
                        if txt_l.value == "-":
                            txt_l.value = "100"
                    page.update()
                return al_cambiar_ol

            dd_opclong.on_select = crear_evento_opclong(dd_opclong, txt_long)

            dd_est = ft.Dropdown(
                options=[ft.dropdown.Option("-"), ft.dropdown.Option("NR"), ft.dropdown.Option("R")],
                value="-", disabled=True, width=140, dense=True, text_size=13, label="Est. posterior"
            )
            dd_tec = ft.Dropdown(
                options=[ft.dropdown.Option("VT"), ft.dropdown.Option("VT/PT"), ft.dropdown.Option("VT/UT")],
                value=tec_def, width=140, dense=True, text_size=13, label="Técnica NDT"
            )
            dd_com = ft.Dropdown(
                options=[ft.dropdown.Option("-"), ft.dropdown.Option("CREAR OT"), ft.dropdown.Option("OT CREADA")],
                value="-", disabled=True, width=180, dense=True, text_size=13, label="Comentarios"
            )

            lbl_pano = ft.Text("Sin foto", size=11, color=ft.Colors.RED_400)
            lbl_det = ft.Text("Sin foto", size=11, color=ft.Colors.RED_400)
            controles[f"img_pano_{key_id}"] = lbl_pano
            controles[f"img_det_{key_id}"] = lbl_det

            contenedor_fotos_rechazo = ft.Container(
                content=ft.Column([
                    ft.Text(f"📷 Fotos de Hallazgo: {desc_z} ({cod_z})", weight="bold", size=12, color=ft.Colors.AMBER_300),
                    ft.Row([
                        ft.Column([boton_foto("Foto Panorámica", f"img_pano_{key_id}"), lbl_pano], spacing=4),
                        ft.Column([boton_foto("Foto de Detalle", f"img_det_{key_id}"), lbl_det], spacing=4),
                    ], wrap=True, spacing=20, run_spacing=10)
                ], spacing=6),
                padding=12, bgcolor=ft.Colors.with_opacity(1, "#161a24"), border_radius=8, visible=False
            )

            def crear_evento_defecto(lbl_c, dd_ol, txt_l, dd_e, dd_cm, box_fotos):
                def al_cambiar_def(e):
                    es_lf = (e.control.value == "LF")
                    if es_lf:
                        lbl_c.value = "ACEPTABLE"
                        lbl_c.color = ft.Colors.GREEN_400
                        dd_ol.visible = False
                        txt_l.visible = True
                        txt_l.value = "-"
                        txt_l.disabled = True
                        dd_e.value = "-"
                        dd_e.disabled = True
                        dd_cm.value = "-"
                        dd_cm.disabled = True
                        box_fotos.visible = False
                    else:
                        lbl_c.value = "RECHAZADO"
                        lbl_c.color = ft.Colors.RED_400
                        dd_ol.visible = True
                        dd_ol.value = "Manual"
                        txt_l.visible = True
                        if txt_l.value == "-":
                            txt_l.value = "100"
                        txt_l.disabled = False
                        dd_e.value = "NR"
                        dd_e.disabled = False
                        dd_cm.value = "CREAR OT"
                        dd_cm.disabled = False
                        box_fotos.visible = True
                    reconstruir_resumen_ots()
                    page.update()
                return al_cambiar_def

            dd_def = ft.Dropdown(
                options=[ft.dropdown.Option(x) for x in ["LF", "D", "DE", "DP", "F", "FA"]],
                value="LF", width=130, dense=True, text_size=13, label="Defecto",
                on_select=crear_evento_defecto(lbl_cond, dd_opclong, txt_long, dd_est, dd_com, contenedor_fotos_rechazo)
            )
            
            controles[f"def_{key_id}"] = dd_def
            controles[f"opclong_{key_id}"] = dd_opclong
            controles[f"longval_{key_id}"] = txt_long
            controles[f"est_{key_id}"] = dd_est
            controles[f"tec_{key_id}"] = dd_tec
            controles[f"com_{key_id}"] = dd_com

            fila_datos = ft.Container(
                content=ft.Column([
                    ft.Row([
                        ft.Text(cod_z, weight="bold", size=15, color=ft.Colors.BLUE_200),
                        ft.Text(desc_z, weight="bold", size=14, expand=True),
                        lbl_cond,
                    ], vertical_alignment=ft.CrossAxisAlignment.CENTER),
                    ft.Text(f"Fecha de inspección: {hoy_str}", size=11, color=ft.Colors.GREY_400),
                    ft.Row([dd_def, col_longitud, dd_est, dd_tec, dd_com], wrap=True, spacing=8, run_spacing=8),
                ], spacing=6),
                padding=10, bgcolor=ft.Colors.with_opacity(1, "#121722"), border_radius=8,
                border=ft.Border.all(1, ft.Colors.BLUE_GREY_800),
            )

            page.add(ft.Column([fila_datos, contenedor_fotos_rechazo], spacing=4))
            page.add(ft.Divider(height=4, color=ft.Colors.TRANSPARENT))
            
        # --- BLOQUES DE LETREROS OBLIGATORIOS (SOLO FOTO PANORÁMICA, FIEL AL ORIGINAL) ---
        if idx_z == 1:
            lbl_pano_rh = ft.Text("Sin foto", size=11, color=ft.Colors.RED_400)
            lbl_pano_lh = ft.Text("Sin foto", size=11, color=ft.Colors.RED_400)
            controles["img_pano_esp_z2_rh"] = lbl_pano_rh
            controles["img_pano_esp_z2_lh"] = lbl_pano_lh

            page.add(
                ft.Container(
                    content=ft.Column([
                        ft.Text("📷 Letreros Obligatorios Especiales: Laterales RH / LH", weight="bold", color=ft.Colors.YELLOW_600, size=14),
                        ft.Row([
                            ft.Column([boton_foto("Letrero Lateral RH (Panorámico)", "img_pano_esp_z2_rh"), lbl_pano_rh], spacing=4),
                            ft.Column([boton_foto("Letrero Lateral LH (Panorámico)", "img_pano_esp_z2_lh"), lbl_pano_lh], spacing=4),
                        ], wrap=True, spacing=20, run_spacing=10)
                    ], spacing=8),
                    padding=15, bgcolor=ft.Colors.BLUE_GREY_900, border_radius=8
                )
            )
        elif idx_z == 5:
            lbl_pano_post = ft.Text("Sin foto", size=11, color=ft.Colors.RED_400)
            controles["img_pano_esp_z8_post"] = lbl_pano_post

            page.add(
                ft.Container(
                    content=ft.Column([
                        ft.Text("📷 Letrero Obligatorio Especial: Posterior", weight="bold", color=ft.Colors.YELLOW_600, size=14),
                        ft.Column([
                            ft.Column([boton_foto("Letrero Posterior (Panorámico)", "img_pano_esp_z8_post"), lbl_pano_post], spacing=4),
                            ft.Container(
                                content=ft.Text("ℹ️ Requerido para la zona posterior (Z08). El campo 'Detalle' en el Excel se llenará automáticamente con '-'.", size=12, color=ft.Colors.BLUE_200),
                                bgcolor=ft.Colors.BLUE_900, padding=10, border_radius=6
                            )
                        ], spacing=8)
                    ], spacing=8),
                    padding=15, bgcolor=ft.Colors.BLUE_GREY_900, border_radius=8
                )
            )

        page.add(
            ft.Text("Leyenda Defectos: D: Desprendimiento | DE: Desgaste | DP: Desprendimiento Parcial | F: Fisurado | LF: Libre de fisura (Aceptable)", size=11, color=ft.Colors.GREY_500),
            ft.Text("Leyenda Técnica: NR: No Reparado | R: Reparado | VT: Inspección Visual | LP: Líquidos Penetrantes | UT: Ultrasonido", size=11, color=ft.Colors.GREY_500),
            ft.Divider(height=35, color=ft.Colors.BLUE_GREY_800)
        )

    # ==========================================================================
    # SECCIÓN 4: ZONAS A REPARAR (RESUMEN DE OTs)
    # ==========================================================================
    reconstruir_resumen_ots()
    page.add(
        ft.Text("4. Resumen de Órdenes de Trabajo (Zonas a Reparar)", size=22, weight="bold", color=ft.Colors.WHITE),
        contenedor_resumen_ots,
        ft.Divider(height=30, color=ft.Colors.BLUE_GREY_700)
    )

    # ==========================================================================
    # SECCIÓN 5: FIRMA DEL RESPONSABLE DE LA INSPECCIÓN
    # ==========================================================================
    dd_nombre_firma = ft.Dropdown(
        label="Nombre de quien realiza la inspección:",
        options=[ft.dropdown.Option("Dennis Huamani"), ft.dropdown.Option("[Entrada Manual]")],
        value="Dennis Huamani", dense=True
    )
    txt_nombre_firma_manual = ft.TextField(label="Nombre (manual):", value="Dennis Huamani", visible=False, dense=True)

    def al_cambiar_nombre_firma(e):
        txt_nombre_firma_manual.visible = (dd_nombre_firma.value == "[Entrada Manual]")
        page.update()
    dd_nombre_firma.on_select = al_cambiar_nombre_firma

    txt_fecha_firma = ft.TextField(label="Fecha de firma (DD/MM/AAAA):", value=hoy_str, dense=True)
    lbl_firma_archivo = ft.Text("Sin imagen de firma seleccionada", size=12, color=ft.Colors.GREY_400)
    controles["firma_upload"] = lbl_firma_archivo

    page.add(
        ft.Text("5. Firma del Responsable de la Inspección", size=22, weight="bold", color=ft.Colors.WHITE),
        ft.ResponsiveRow([
            ft.Column([dd_nombre_firma, txt_nombre_firma_manual, txt_fecha_firma], col={"sm": 12, "md": 6}),
            ft.Column([
                ft.Text("Subir imagen de firma (PNG/JPG):", size=13, color=ft.Colors.GREY_300),
                ft.Button("Seleccionar imagen de firma", icon=ft.Icons.UPLOAD_FILE, on_click=lambda e: abrir_selector("firma_upload")),
                lbl_firma_archivo
            ], col={"sm": 12, "md": 6}),
        ]),
        ft.Divider(height=30, color=ft.Colors.BLUE_GREY_700)
    )

    # ==========================================================================
    # SECCIÓN 6: GENERAR REPORTE FINAL
    # ==========================================================================
    texto_estado = ft.Text("", size=16, weight="bold")
    
    async def procesar_todo(e):
        ruta_plantilla = buscar_recurso("plantilla_tolva.xlsx")
        if not ruta_plantilla:
            texto_estado.value = (
                "❌ No se encontró 'plantilla_tolva.xlsx' dentro de la app.\n"
                "Archivos que veo:\n" + diagnostico_archivos()
            )
            texto_estado.color = ft.Colors.RED_400
            page.update()
            return

        texto_estado.value = "⚙️ Generando el archivo Excel..."
        texto_estado.color = ft.Colors.BLUE_300
        page.update()
        
        for key, control in controles.items():
            if isinstance(control, (ft.Dropdown, ft.TextField)):
                estado_app[key] = control.value
                
        todos_los_rechazos = []
        for idx_z, zona in enumerate(ESTRUCTURA_ZONAS):
            for idx_i, item in enumerate(zona["items"]):
                cod_z, desc_z, tec_def = item
                key_id = f"z{idx_z}_{idx_i}"
                def_val = estado_app.get(f"def_{key_id}", "LF")
                if def_val != "LF":
                    opc_l = estado_app.get(f"opclong_{key_id}", "Manual")
                    long_final = "VARIOS" if opc_l == "VARIOS" else estado_app.get(f"longval_{key_id}", "100")
                    todos_los_rechazos.append({
                        "zona": cod_z,
                        "descripcion": desc_z,
                        "defecto": def_val,
                        "longitud": long_final,
                        "est_post": estado_app.get(f"est_{key_id}", "NR"),
                        "tecnica": estado_app.get(f"tec_{key_id}", tec_def),
                        "comentario": estado_app.get(f"com_{key_id}", "CREAR OT"),
                        "key_id": key_id
                    })
        
        try:
            cliente_final = txt_cliente_manual.value if dd_cliente.value == "[Entrada Manual]" else dd_cliente.value
            lugar_final = txt_lugar_manual.value if dd_lugar.value == "[Entrada Manual]" else dd_lugar.value
            pm_final = obtener_pm_efectivo()
            nombre_firma_final = txt_nombre_firma_manual.value if dd_nombre_firma.value == "[Entrada Manual]" else dd_nombre_firma.value
            cod_inf_final = f"ZA-IF-{dd_equipo.value}-{txt_sufijo.value.strip()}"

            try:
                horometro_val = float(txt_horometro.value.strip())
            except Exception:
                horometro_val = txt_horometro.value.strip()

            # Construir el DataFrame 8x7 de espesores si el PM corresponde
            df_matriz_final = None
            if es_pm_espesores(pm_final):
                if chk_no_medicion.value:
                    df_matriz_final = pd.DataFrame(
                        [["-"] * 7 for _ in range(8)],
                        index=[f"Punto {i+1}" for i in range(8)],
                        columns=[f"Eje {j+1}" for j in range(7)]
                    )
                else:
                    datos_m = []
                    for i in range(8):
                        fila_m = []
                        for j in range(7):
                            val_celda = celdas_matriz[i][j].value.strip()
                            try:
                                fila_m.append(round(float(val_celda), 2))
                            except Exception:
                                fila_m.append(val_celda if val_celda else "-")
                        datos_m.append(fila_m)
                    df_matriz_final = pd.DataFrame(
                        datos_m,
                        index=[f"Punto {i+1}" for i in range(8)],
                        columns=[f"Eje {j+1}" for j in range(7)]
                    )

            excel_bytes, avisos = await asyncio.to_thread(
                lambda: generar_reporte_excel(
                ruta_plantilla=ruta_plantilla,
                cliente=cliente_final,
                lugar=lugar_final,
                fecha_insp=txt_fecha_insp.value.strip(),
                cod_equipo=dd_equipo.value,
                cod_tolva=txt_tolva.value.strip(),
                horometro=horometro_val,
                cod_informe=cod_inf_final,
                revision=txt_revision.value.strip(),
                pm=pm_final,
                estructura_zonas=ESTRUCTURA_ZONAS,
                nombre_realizado=nombre_firma_final,
                fecha_firma=txt_fecha_firma.value.strip(),
                firma_ruta=rutas_fotos.get("firma_upload"),
                todos_los_rechazos=todos_los_rechazos,
                matriz_espesores=df_matriz_final
                )
            )
            
            fecha_limpia = txt_fecha_insp.value.strip().replace("/", "")
            nombre_salida = f"{cod_inf_final}_{fecha_limpia}.xlsx"
            texto_estado.value = "💾 Elige dónde guardar el Excel..."
            page.update()
            ruta_guardada = await selector_fotos.save_file(
                dialog_title="Guardar reporte Excel",
                file_name=nombre_salida,
                file_type=ft.FilePickerFileType.CUSTOM,
                allowed_extensions=["xlsx"],
                src_bytes=excel_bytes,
            )
            if not page.web and not es_movil():
                # En PC el selector solo devuelve la ruta: escribimos el archivo nosotros.
                ruta_guardada = ruta_guardada or nombre_salida
                with open(ruta_guardada, "wb") as f:
                    f.write(excel_bytes)
            if not ruta_guardada and not page.web:
                texto_estado.value = "⚠️ Guardado cancelado. Vuelve a pulsar el botón para elegir dónde guardarlo."
                texto_estado.color = ft.Colors.AMBER_400
                page.update()
                return

            msg = f"✅ Reporte Excel generado correctamente: {ruta_guardada or nombre_salida}"
            if avisos:
                msg += "\n" + "\n".join(avisos)
            texto_estado.value = msg
            texto_estado.color = ft.Colors.GREEN_400
        except Exception as ex:
            texto_estado.value = f"❌ Error al generar el reporte: {ex}"
            texto_estado.color = ft.Colors.RED_400
        page.update()

    page.add(
        ft.Text("6. Generar Reporte Final", size=22, weight="bold", color=ft.Colors.WHITE),
        ft.Container(
            content=ft.Column([
                ft.Button("📥 GENERAR REPORTE EXCEL", on_click=procesar_todo, bgcolor=ft.Colors.BLUE_700, color=ft.Colors.WHITE, height=55, width=380),
                texto_estado
            ], horizontal_alignment=ft.CrossAxisAlignment.CENTER),
            alignment=ft.Alignment(0, 0),
            margin=ft.Margin(left=0, top=15, right=0, bottom=50)
        )
    )

ft.run(main)